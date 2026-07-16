from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUTPUT = Path(__file__).with_name("公路工程智慧试验室机器人自动化作业技术规程_完整初稿.docx")


def set_run_font(run, cn_font="宋体", en_font="Times New Roman", size=10.5, bold=False):
    run.font.name = en_font
    run.font.size = Pt(size)
    run.font.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), cn_font)
    r_fonts.set(qn("w:ascii"), en_font)
    r_fonts.set(qn("w:hAnsi"), en_font)


def set_paragraph_format(paragraph, first_indent=True, before=0, after=0, line=1.5):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.line_spacing = line
    if first_indent:
        fmt.first_line_indent = Cm(0.74)
    else:
        fmt.first_line_indent = Cm(0)


def add_para(doc: Document, text: str = "", first_indent=True, bold_prefix: str | None = None,
             align=None, size=10.5, cn_font="宋体", after=0):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_indent=first_indent, after=after)
    if align is not None:
        p.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, cn_font=cn_font, size=size, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, cn_font=cn_font, size=size)
    else:
        r = p.add_run(text)
        set_run_font(r, cn_font=cn_font, size=size)
    return p


def add_clause(doc: Document, number: str, text: str):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_indent=False)
    r1 = p.add_run(f"{number} ")
    set_run_font(r1, size=10.5, bold=False)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5)
    return p


def add_list_item(doc: Document, marker: str, text: str):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_indent=False)
    p.paragraph_format.left_indent = Cm(0.74)
    p.paragraph_format.hanging_indent = Cm(0.74)
    r1 = p.add_run(f"{marker} ")
    set_run_font(r1, size=10.5)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5)
    return p


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_paragraph()
    p.style = doc.styles[f"Heading {level}"]
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run(text)
    if level == 1:
        set_run_font(r, cn_font="黑体", size=15, bold=True)
    elif level == 2:
        set_run_font(r, cn_font="黑体", size=12, bold=True)
    else:
        set_run_font(r, cn_font="黑体", size=10.5, bold=True)
    return p


def set_cell_text(cell, text: str, bold=False, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(str(text))
    set_run_font(r, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade_cell(cell, fill="D9EAF7"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_table(doc: Document, headers: Sequence[str], rows: Iterable[Sequence[str]], widths=None, font_size=8.5):
    rows = list(rows)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=font_size)
        shade_cell(table.rows[0].cells[i])
        if widths:
            table.rows[0].cells[i].width = Cm(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=font_size)
            if widths:
                cells[i].width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=9)


def add_toc(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("目  次")
    set_run_font(r, cn_font="黑体", size=18, bold=True)
    p.paragraph_format.space_after = Pt(18)

    toc_p = doc.add_paragraph()
    toc_p.paragraph_format.first_line_indent = Cm(0)
    run = toc_p.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "在 Word 中右键选择“更新域”以更新目录。"
    fld_char3 = OxmlElement("w:fldChar")
    fld_char3.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    run._r.append(text)
    run._r.append(fld_char3)
    doc.add_page_break()


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.4)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    for level, cn_font, size in [(1, "黑体", 15), (2, "黑体", 12), (3, "黑体", 10.5)]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style._element.rPr.rFonts.set(qn("w:eastAsia"), cn_font)

    settings = doc.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def add_cover(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("ICS 93.080\nCCS P 66\nT/XXXX XXX—202X")
    set_run_font(r, size=11)

    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("团  体  标  准")
    set_run_font(r, cn_font="黑体", size=24, bold=True)

    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("公路工程智慧试验室机器人\n自动化作业技术规程")
    set_run_font(r, cn_font="黑体", size=26, bold=True)
    p.paragraph_format.line_spacing = 1.2

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Technical specification for robotic automated operations in intelligent laboratories for highway engineering")
    set_run_font(r, size=12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("（征求意见稿·初稿）")
    set_run_font(r, cn_font="宋体", size=14)

    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("202X-XX-XX 发布                         202X-XX-XX 实施")
    set_run_font(r, size=11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("XXXX 协会（学会）  发布")
    set_run_font(r, cn_font="黑体", size=16, bold=True)
    doc.add_page_break()


def add_front_matter(doc: Document):
    add_heading(doc, "前言", 1)
    add_para(doc, "本文件按照 GB/T 1.1—2020《标准化工作导则  第1部分：标准化文件的结构和起草规则》的规定起草。")
    add_para(doc, "请注意本文件的某些内容可能涉及专利。本文件的发布机构不承担识别专利的责任。")
    add_para(doc, "本文件由 XXXX 提出。")
    add_para(doc, "本文件由 XXXX 归口。")
    add_para(doc, "本文件起草单位：XXXX、XXXX、XXXX。")
    add_para(doc, "本文件主要起草人：XXXX、XXXX、XXXX。")
    add_para(doc, "本文件为首次发布。")
    doc.add_page_break()

    add_heading(doc, "引言", 1)
    add_para(doc, "公路工程工地试验室承担原材料进场检验、配合比验证、施工过程质量控制、试件制作与养护、试验检测以及质量数据归档等任务。随着施工规模扩大、试验频次提高和质量追溯要求增强，样品搬运、称量投料、制样成型、仪器上下料、筛分称量、清洁换样和废样处置等高频重复作业对自动化、标准化和安全性的需求日益突出。")
    add_para(doc, "机器人自动化作业能够降低重复劳动和高温、粉尘、夹伤、飞溅等作业风险，但其应用涉及样品身份、试验方法、动作程序、试验仪器、末端工具、安全联锁和数据系统等多环节协同。若缺少统一要求，可能发生样品错配、作业参数失控、交叉污染、原始数据错关联、异常处置不完整以及自动作业影响试验结果等问题。")
    add_para(doc, "本文件以样品唯一身份为主索引，以试验任务配方为作业依据，以机器人、末端执行器、试验仪器和信息系统协同为基础，建立“任务核验—自动作业—结果关联—人工审核—异常处置—质量追溯”的全过程控制机制。本文件规定机器人自动化作业的实现和验收要求，不改变现行试验方法规定的取样、制样、试验、计算和结果判定要求。")
    doc.add_page_break()


def chapter_1_4(doc: Document):
    add_heading(doc, "1 范围", 1)
    add_para(doc, "本文件规定了公路工程智慧试验室机器人自动化作业的基本要求、系统组成与功能、试验室布置与设备、样品与试件管理、自动化作业、数据与接口、异常处置、人工接管、安全防护、安装调试、作业验证、系统检验与验收、运行维护以及数据安全要求。")
    add_para(doc, "本文件适用于高速公路及其他公路工程施工项目工地试验室内样品识别与流转、称量取料与投料、辅助制样、试件养护流转、试验仪器上下料、筛分分选、清洁换样和废样处置等机器人自动化作业。中心试验室、第三方试验检测机构和其他工程材料试验室可参照执行。")
    add_para(doc, "本文件不替代现行标准规定的取样、制样、试验、计算、结果审核、质量判定和试验检测人员执业要求。自动化作业尚未通过验证或有效监测条件不满足时，应采用经批准的人工或替代作业方式。")

    add_heading(doc, "2 规范性引用文件", 1)
    add_para(doc, "下列文件中的内容通过文中的规范性引用而构成本文件必不可少的条款。其中，注日期的引用文件，仅该日期对应的版本适用于本文件；不注日期的引用文件，其最新版本（包括所有的修改单）适用于本文件。")
    refs = [
        "GB/T 4208  外壳防护等级（IP代码）",
        "GB/T 5226.1  机械电气安全  机械电气设备  第1部分：通用技术条件",
        "GB 11291.1  工业环境用机器人  安全要求  第1部分：机器人",
        "GB 11291.2  机器人与机器人装备  工业机器人的安全要求  第2部分：机器人系统与集成",
        "GB/T 12643  机器人与机器人装备  词汇",
        "GB/T 15706  机械安全  设计通则  风险评估与风险减小",
        "GB/T 16754  机械安全  急停功能  设计原则",
        "GB/T 16855.1  机械安全  控制系统安全相关部件  第1部分：设计通则",
        "GB/T 22239  信息安全技术  网络安全等级保护基本要求",
        "GB/T 36344  信息技术  数据质量评价指标",
        "GB/T 38155  重要产品追溯  追溯术语",
        "GB/T 41255  智能工厂  通用技术要求",
        "GB/T 42982  工业机器人平均无故障工作时间计算方法",
        "GB/T 42983（所有部分）  工业机器人  运行维护",
        "GB/Z 43065.1  机器人  工业机器人系统的安全设计  第1部分：末端执行器",
        "GB/Z 43065.2  机器人  工业机器人系统的安全设计  第2部分：手动装载/卸载工作站",
        "JTG 3420  公路工程水泥及水泥混凝土试验规程",
        "JTG 3430  公路土工试验规程",
        "JTG 3432  公路工程集料试验规程",
        "JTG E20  公路工程沥青及沥青混合料试验规程",
        "JTG E51  公路工程无机结合料稳定材料试验规程",
        "JTG 3450  公路路基路面现场测试规程"
    ]
    for ref in refs:
        add_para(doc, ref, first_indent=False)

    add_heading(doc, "3 术语和定义", 1)
    terms = [
        ("3.1", "智慧试验室 intelligent laboratory", "利用自动化装备、感知设备、机器人、试验仪器和信息系统，对样品、作业、数据和质量活动进行协同管理的试验室。"),
        ("3.2", "机器人自动化作业系统 robotic automated operation system", "由机器人本体、末端执行器、试验仪器、辅助装置、控制系统、安全防护和信息系统组成，用于自动完成试验室作业任务的系统。"),
        ("3.3", "机器人作业单元 robotic work cell", "在限定工作空间内，由机器人及其配套设备共同完成一类或多类自动化作业的基本功能单元。"),
        ("3.4", "末端执行器 end effector", "安装于机器人机械接口、用于抓取、吸附、取料、投料、清洁、操作仪器或执行其他任务的装置。"),
        ("3.5", "试验任务配方 test task recipe", "与特定试验项目和作业对象对应的样品要求、动作顺序、工具、仪器参数、环境条件、判定规则和异常处理规则的集合。"),
        ("3.6", "样品身份链 sample identity chain", "样品从接收、拆分、制样、养护、试验、留样至处置全过程的唯一身份及其关联关系。"),
        ("3.7", "作业关键点 operation critical point", "影响样品身份、试样状态、试验条件、人员安全或结果可靠性的机器人动作、设备状态或数据节点。"),
        ("3.8", "作业成功率 operation success rate", "统计周期内无需非计划人工干预且满足任务完成条件的自动作业任务数与应执行任务总数的比值。"),
        ("3.9", "人工接管 manual takeover", "自动作业无法安全、正确继续时，由授权人员暂停、调整、恢复或终止任务的过程。"),
        ("3.10", "一样一档 one file for one sample", "以样品或试件唯一编码为索引，汇集其委托、流转、作业、试验、异常、审核和处置信息形成的数字化档案。"),
        ("3.11", "交叉污染 cross-contamination", "前序样品、工具、容器、设备或环境中的残留物对后续样品状态或试验结果造成影响的现象。"),
        ("3.12", "安全工作空间 safeguarded space", "通过固定或活动防护、感知装置和安全控制限制人员进入的机器人危险运动空间。")
    ]
    for num, term, definition in terms:
        add_heading(doc, f"{num} {term}", 2)
        add_para(doc, definition)

    add_heading(doc, "4 基本规定", 1)
    clauses = [
        ("4.1", "机器人自动化作业应纳入试验室质量管理体系，并与委托管理、样品管理、仪器设备管理、试验任务、原始记录、报告审核和不合格结果管理协同实施。"),
        ("4.2", "自动化作业应遵循试验方法优先、样品身份唯一、动作过程可验证、原始数据可追溯、安全联锁优先和异常可人工接管的原则。"),
        ("4.3", "机器人自动化作业不得改变试样质量、温度、湿度、几何形状、受力状态、养护龄期以及现行试验方法规定的其他条件。"),
        ("4.4", "机器人承担的作业环节、适用样品、试验方法、设备、工具、程序和限制条件应在试验任务配方中明确。未经验证的程序、工具或任务配方不得用于正式试验。"),
        ("4.5", "每件样品、试件、容器、工具、机器人、试验仪器和工位应具有唯一编码。样品与委托单、试验任务、试验方法、设备和结果之间的关联应可追溯。"),
        ("4.6", "试验仪器应按计量和试验方法要求检定、校准或校验。机器人系统不得绕过仪器安全条件、校准状态和试验启动条件直接执行试验。"),
        ("4.7", "用于正式试验的自动化作业应通过现场作业验证和人工平行比对。验证结果不满足第12章和第13章要求时，可保留辅助搬运或信息提示功能，不得自动完成影响试验结果的关键操作。"),
        ("4.8", "机器人系统不得修改试验仪器产生的原始数据。数据补录、修正和重新关联应保留原值、修正值、原因、人员和时间。"),
        ("4.9", "安全控制功能应独立于一般业务控制功能。急停、防护门、区域入侵、工具锁紧和危险设备联锁不得仅依赖上位机软件或普通网络通信实现。"),
        ("4.10", "自动化作业出现样品身份不明、工具状态不明、设备校准失效、安全条件不满足、数据关联失败或试样状态可能受到影响时，应阻断任务并生成异常记录。"),
        ("4.11", "机器人自动化作业不得削减现行标准和质量管理体系规定的见证、审核、复核及人员签字要求。"),
        ("4.12", "新建智慧试验室宜将机器人系统与试验室布局、样品流线、仪器接口、供电通信和安全防护同步设计、同步安装、同步调试和同步验收。"),
        ("4.13", "既有试验室改造期间应保证正常试验和数据连续，改造作业不得污染样品、影响仪器稳定性或形成新的安全风险。")
    ]
    for num, text in clauses:
        add_clause(doc, num, text)


def chapter_5_8(doc: Document):
    add_heading(doc, "5 系统组成与功能", 1)
    add_heading(doc, "5.1 系统组成", 2)
    add_clause(doc, "5.1.1", "机器人自动化作业系统宜由样品与任务管理、机器人与运动控制、末端执行器与工具管理、试验仪器与辅助设备、视觉与识别、安全控制、数据与接口、应用与追溯等部分组成。")
    add_clause(doc, "5.1.2", "系统架构应区分业务管理、作业控制和安全控制。涉及人员防护和危险运动停止的功能应由符合安全要求的控制回路实现。")
    add_clause(doc, "5.1.3", "单个机器人作业单元可完成一种或多种试验室作业，但各任务的工具、程序、工位、样品和数据应相互隔离，避免错用和交叉污染。")

    add_heading(doc, "5.2 样品与任务管理功能", 2)
    for num, text in [
        ("5.2.1", "系统应接收委托信息、样品信息、试验项目、试验方法、计划时间和优先级，并生成唯一试验任务。"),
        ("5.2.2", "系统应在任务开始前核验样品身份、试验方法、仪器校准状态、工具状态、耗材、环境条件和上一任务清洁状态。"),
        ("5.2.3", "系统应支持任务暂停、恢复、取消、重新执行和人工转办，并记录状态变化原因。"),
        ("5.2.4", "同一样品拆分形成多个试样或试件时，应保留父子关联；多个分样合并使用时，应记录来源和合并规则。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "5.3 机器人与工具控制功能", 2)
    for num, text in [
        ("5.3.1", "系统应具备程序调用、坐标转换、工具识别、抓取确认、路径控制、到位确认、碰撞监测和状态记录功能。"),
        ("5.3.2", "采用自动换装时，应在机器人运动前确认工具编码、机械锁紧、能源连接和状态反馈一致。"),
        ("5.3.3", "涉及称量、投料、试件对中和精细操作时，应保存关键动作位置、姿态、速度、力或其他可验证参数。"),
        ("5.3.4", "任务中断后恢复运行前，应重新核验机器人位置、工具、样品、工位和仪器状态，不得仅从程序断点直接继续。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "5.4 试验仪器协同功能", 2)
    for num, text in [
        ("5.4.1", "系统应通过可验证接口获取试验仪器准备、运行、完成、故障、防护和数据状态。"),
        ("5.4.2", "仪器不满足校准有效、空载复位、防护关闭、样品到位或环境条件时，机器人不得自动启动试验。"),
        ("5.4.3", "机器人和试验仪器的控制权限、启动条件、停止条件和故障责任边界应在接口文件中明确。"),
        ("5.4.4", "接口通信失败时，应保持设备处于安全状态，并防止重复启动、重复记账和数据覆盖。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "5.5 数据与应用功能", 2)
    for num, text in [
        ("5.5.1", "系统应记录样品、任务、机器人、工具、仪器、程序、时间、动作结果、原始数据、异常和人工操作。核心数据字段应符合附录B的规定。"),
        ("5.5.2", "系统应支持按样品、试验项目、任务、设备、时间和状态检索，并可由结果追溯至原始记录和作业证据。"),
        ("5.5.3", "系统应具备作业统计、异常统计、任务成功率、设备在线率、人工接管次数和清洁状态分析功能。"),
        ("5.5.4", "系统应支持形成一样一档，并与试验室信息管理系统交换委托、样品、任务、结果和审核状态。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "6 试验室布置与设备要求", 1)
    add_heading(doc, "6.1 试验室布置", 2)
    clauses = [
        ("6.1.1", "试验室布置应根据样品流、人员流、工具流和废弃物流进行分区，宜设置收样与暂存区、称量取料区、制样区、养护区、试验区、清洁区和废样区。"),
        ("6.1.2", "机器人工作空间、人员通道和设备维修空间应明确标识。非协作状态下，人员通道不得穿越机器人危险运动空间。"),
        ("6.1.3", "称量和精密测量区域应与强振动、气流、粉尘、清洗和废样破碎区域隔离。"),
        ("6.1.4", "压力试验、抗折试验及其他可能产生碎块飞溅的工位应设置防护装置，机器人和人员的安全位置应满足试验设备防护要求。"),
        ("6.1.5", "高温烘箱、加热设备、沥青及其他高温材料作业区应设置隔热、防烫、通风和异常温度监测措施。"),
        ("6.1.6", "清洁区应具备与材料相适应的集尘、给排水、废液和固体废弃物收集设施，不得使清洗介质进入称量、养护和试验区域。"),
        ("6.1.7", "机器人基础、轨道、货架和工位应具有足够的刚度、稳定性和承载能力，不得因地面沉降、设备振动或货架变形影响作业精度。")
    ]
    for num, text in clauses:
        add_clause(doc, num, text)

    add_heading(doc, "6.2 机器人本体", 2)
    for num, text in [
        ("6.2.1", "机器人额定负载、工作范围、自由度、速度、重复定位能力和环境适应性应满足样品、工具和作业工位要求，并留有负载和工作空间余量。"),
        ("6.2.2", "机器人在额定负载、最大允许偏心距和实际工具惯量条件下应保持稳定运行。工具、样品和附件总质量不得超过允许负载。"),
        ("6.2.3", "用于量化定位、投料、对中或仪器操作的机器人，其综合位置和姿态误差宜不大于相应作业允许偏差绝对值的三分之一。"),
        ("6.2.4", "机器人应具备断电制动或安全保持措施。断电、急停或通信故障时，不得造成样品、工具或载荷失控坠落。"),
        ("6.2.5", "机器人外壳防护等级应根据粉尘、潮湿、腐蚀和清洗环境确定，并符合GB/T 4208的规定。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "6.3 末端执行器与工具", 2)
    for num, text in [
        ("6.3.1", "末端执行器应根据样品材质、形状、质量、温度、表面状态和试验方法选用，夹持或吸附不得造成破损、掉角、变形、污染或状态改变。"),
        ("6.3.2", "抓取装置应具备夹持、吸附或锁紧状态反馈。对坠落可能造成伤害、样品失效或设备损坏的任务，应设置冗余保持或失效安全措施。"),
        ("6.3.3", "与样品接触的工具和容器材料应与试验对象相容，易粘附、易吸水、易腐蚀或高温作业时应采取专用工具。"),
        ("6.3.4", "自动换装接口应具备防错连接和机械锁紧确认功能。工具识别及锁紧确认率应达到100%。"),
        ("6.3.5", "工具应设置唯一编码和使用状态。发生碰撞、松动、磨损、变形、吸附能力下降或清洁不合格时应停用。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "6.4 试验仪器与辅助设备", 2)
    for num, text in [
        ("6.4.1", "天平、烘箱、养护设施、振筛机、压力试验机、搅拌和成型设备等应保留原有计量、控制和安全功能。机器人接入不得降低设备安全等级和计量性能。"),
        ("6.4.2", "设备接口应说明信号名称、数据类型、单位、精度、时间基准、状态码、控制权限、异常码和版本号。"),
        ("6.4.3", "机器人上下料工位应设置样品到位、方向、数量和防护状态检测。仅依赖预设坐标且无到位确认的方式不得用于关键试验操作。"),
        ("6.4.4", "用于样品储存、养护和冷却的货架或托盘应具备位置编码、防错放和承载状态确认。"),
        ("6.4.5", "视觉识别用于尺寸、位置或状态判定时，应设置稳定的光照、背景和标定基准；视场、焦距或安装位置变化后应重新标定。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "6.5 环境条件", 2)
    for num, text in [
        ("6.5.1", "作业环境的温度、湿度、振动、粉尘和照明应同时满足机器人、试验仪器和相应试验方法要求。"),
        ("6.5.2", "系统应监测影响自动作业和试验结果的环境条件。环境超出任务配方允许范围时，应阻断相关任务或转为人工确认。"),
        ("6.5.3", "压缩空气、真空、供水、排水和电源中断时，应使机器人和样品保持安全，并生成设备异常。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "7 样品与试件管理", 1)
    add_heading(doc, "7.1 编码与身份核验", 2)
    for num, text in [
        ("7.1.1", "样品接收后应赋予唯一编码。需要拆分、制样或成型的，应为子样、试样、试件和容器建立唯一编码，并保留来源关系。"),
        ("7.1.2", "样品编码载体应与环境相适应，可采用二维码、射频标签、耐久字符或其他方式。编码载体不得影响样品和试验结果。"),
        ("7.1.3", "机器人抓取前和试验仪器接收前应进行身份核验。无法识读、重复编码、身份冲突或样品与任务不一致时，不得继续自动作业。"),
        ("7.1.4", "样品、试验任务和试验设备的关联准确率应达到100%。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "7.2 接收、储存与养护", 2)
    for num, text in [
        ("7.2.1", "样品接收时应记录名称、规格、批次、数量、状态、来源、接收时间和试验要求。发现破损、泄漏、污染或标识不清时应人工确认。"),
        ("7.2.2", "储存和养护位置应唯一编码，并记录入库、出库、移动和占用状态。机器人不得将样品放入与材料、温度、湿度或龄期要求不相容的位置。"),
        ("7.2.3", "具有龄期、温度或时间要求的试件，系统应以统一时间基准计算可试验时间，未达到条件或超过允许时间窗时应阻断试验并提示复核。"),
        ("7.2.4", "样品堆放和托盘布置应防止倾倒、挤压和混样。机器人取放顺序不得影响其他样品稳定性。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "7.3 样品流转与状态", 2)
    for num, text in [
        ("7.3.1", "样品流转应记录起始工位、目标工位、机器人、工具、开始时间、完成时间和结果。"),
        ("7.3.2", "样品状态宜包括待接收、待制样、制样中、待养护、养护中、待试验、试验中、待审核、留样、废样和异常。"),
        ("7.3.3", "自动作业造成样品掉落、破损、表面污染、温度异常或状态不确定时，应将样品标记为异常并隔离，未经授权人员确认不得重新投入试验。"),
        ("7.3.4", "留样和废样处置应执行批准的期限、位置和方式，处置前应核验样品身份和审核状态。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "8 机器人自动化作业", 1)
    add_heading(doc, "8.1 一般规定", 2)
    for num, text in [
        ("8.1.1", "自动化作业应按试验任务配方执行。任务配方应至少包括适用样品、试验方法、作业流程、设备工具、关键参数、完成条件、异常规则和人工接管点。"),
        ("8.1.2", "影响试验结果的关键动作应设置到位、数量、质量、时间或状态确认。未获得有效确认时，不得进入下一作业步骤。"),
        ("8.1.3", "任务配方、机器人程序、仪器方法和判定规则应采用受控版本。正式运行期间不得未经授权在线修改。"),
        ("8.1.4", "自动化作业的典型项目和最低功能应符合附录A的规定。未列入附录A的新作业应按第12章完成验证后应用。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "8.2 任务接收与作业准备", 2)
    for num, text in [
        ("8.2.1", "任务开始前应核验样品身份、任务状态、试验方法、仪器校准有效期、工具、耗材、环境、安全防护和清洁状态。"),
        ("8.2.2", "系统应防止同时向同一样品、仪器或工位下发冲突任务。任务队列调整应保留原因和操作记录。"),
        ("8.2.3", "机器人启动前应确认起始位置、工具中心点、工件坐标系和工作空间无异常。发生重新上电、人工移动、碰撞或工具更换后应重新确认。"),
        ("8.2.4", "涉及高温、压力、旋转、切割或破坏性试验的设备应在任务开始前完成安全联锁自检。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "8.3 抓取、搬运与定位", 2)
    for num, text in [
        ("8.3.1", "机器人应依据样品形状、质量和表面状态选择抓取位置、姿态和夹持方式。抓取后应确认样品数量、身份和保持状态。"),
        ("8.3.2", "搬运路径应避开人员、仪器、防护设施和其他样品。速度和加速度不得造成样品飞散、液体溢出、粉料扬尘或试件碰撞。"),
        ("8.3.3", "样品放置后应确认位置、方向、支承和稳定性。对中、插入、装夹等精细操作应具备视觉、力觉或机械导向等纠偏措施。"),
        ("8.3.4", "抓取成功率应不低于99%；连续发生2次抓取或放置失败时，应暂停任务并转入人工接管。"),
        ("8.3.5", "机器人不得以碰撞限位代替正常到位判定。发生非预期接触时应记录位置、载荷和任务状态。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "8.4 称量、取料与投料", 2)
    for num, text in [
        ("8.4.1", "称量前应确认天平校准、水平、去皮和稳定状态。样品、容器和称量结果应自动关联。"),
        ("8.4.2", "机器人放置称量容器时不得碰撞称量盘或形成偏载。称量值未稳定或超出量程时不得记录为有效结果。"),
        ("8.4.3", "自动取料和投料应按任务配方控制目标质量、允许偏差、投料顺序和残留。允许偏差应符合相应试验方法；无明确规定时应通过比对试验确定。"),
        ("8.4.4", "分次投料时应保存每次实测值和累计值，不得仅保存最终汇总结果。发生撒漏、粘附或残留超限时应生成异常并重新确认质量。"),
        ("8.4.5", "不同材料、批次或试验项目切换前应确认容器和工具清洁状态，防止交叉污染。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "8.5 制样、成型与养护流转", 2)
    for num, text in [
        ("8.5.1", "自动制样应按现行试验方法控制材料质量、投料顺序、搅拌时间、成型方式、压实或振实参数、试模状态和环境条件。"),
        ("8.5.2", "机器人参与装模、刮平、插捣、振实或脱模时，不得引入超出试验方法允许的离析、损伤、气泡、尺寸偏差或表面缺陷。"),
        ("8.5.3", "试件成型后应自动关联样品、试模、成型时间、养护条件和计划试验龄期。标识应在搬运、养护和试验过程中保持可识读。"),
        ("8.5.4", "试件入养护设施前应确认外观、数量、位置和养护设备状态；出库时应核验龄期、库位、温湿度记录和试验任务。"),
        ("8.5.5", "脱模和搬运过程中出现掉角、裂纹、变形或标识损坏时，应隔离试件并由授权人员判定其有效性。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "8.6 仪器上下料与试验操作", 2)
    for num, text in [
        ("8.6.1", "机器人向试验仪器装载样品前，应核验仪器方法、量程、校准状态、空载状态、防护状态和样品身份。"),
        ("8.6.2", "试件或样品应按试验方法规定的位置和方向放置。采用自动对中时，应验证对中误差不影响试验结果。"),
        ("8.6.3", "仪器启动指令应在样品到位、防护关闭、机器人退出危险区和其他启动条件全部满足后发出。"),
        ("8.6.4", "试验过程中机器人不得进入仪器危险运动区域。需要协同动作时，应采用经风险评估和验证的安全控制方式。"),
        ("8.6.5", "试验完成后，应确认仪器卸载、样品可安全移出和原始数据已生成。破坏性试验的残片应按规定收集，不得与未试验样品混放。"),
        ("8.6.6", "试验设备异常停机、数据曲线不完整或样品位置变化时，不得自动判定任务完成，应转入人工复核。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "8.7 筛分、分选与结果采集", 2)
    for num, text in [
        ("8.7.1", "自动筛分前应核验筛具规格、筛孔顺序、清洁状态、底盘和盖板。筛组与试验任务的匹配正确率应达到100%。"),
        ("8.7.2", "样品投放、筛分时间、振筛参数和人工辅助筛分应符合相应试验方法。自动系统不得以设备运行结束代替筛分完成判定。"),
        ("8.7.3", "各级筛余应分别称量并与对应筛孔自动关联。筛余总质量与筛前质量的差值应符合试验方法要求，超限时应检查遗撒、残留和错放。"),
        ("8.7.4", "系统宜识别筛孔堵塞、筛具破损、样品团聚和异常残留；无法自动确认时应提示人工检查。"),
        ("8.7.5", "自动分选产生的各级样品应使用独立容器，并保持编码、质量和去向可追溯。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "8.8 清洁、换样与废弃物处置", 2)
    for num, text in [
        ("8.8.1", "自动作业系统应根据材料特性和试验项目制定清洁配方，明确清洁对象、介质、方法、时间、完成条件和废弃物去向。"),
        ("8.8.2", "不同材料、不同批次或可能相互影响的试验项目切换前，应完成工具、容器、工位和设备接触面的清洁确认。"),
        ("8.8.3", "清洁完成状态可通过质量差、视觉、压力、流量、时间或人工确认判定。未经确认不得开始下一任务。"),
        ("8.8.4", "粉尘、废液、高温残料、破坏试件和锐利碎片应分类收集。机器人转运不得造成扬尘、泄漏、飞溅和二次污染。"),
        ("8.8.5", "废样处置前应确认试验完成、结果审核和留样期限。误处置、提前处置或身份不明时应生成关键异常。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "8.9 作业完成与结果确认", 2)
    for num, text in [
        ("8.9.1", "自动任务完成条件应至少包括规定动作完成、样品状态正确、仪器数据生成、异常已处理和作业记录完整。"),
        ("8.9.2", "系统应区分自动采集值、计算值、机器人判定、仪器判定和人工审核结论。自动任务完成不等同于试验结果审核通过。"),
        ("8.9.3", "未完成、人工接管或异常恢复的任务应在结果页面和一样一档中明确标识，不得按正常全自动任务统计。"),
        ("8.9.4", "自动作业任务成功率应按附录B计算，正式验收和正常运行期间均不应低于98%。")
    ]:
        add_clause(doc, num, text)


def chapter_9_14(doc: Document):
    add_heading(doc, "9 数据、接口与质量追溯", 1)
    add_heading(doc, "9.1 数据采集与编码", 2)
    for num, text in [
        ("9.1.1", "系统数据宜分为基础数据、样品数据、任务数据、动作数据、仪器数据、环境数据、异常数据、审核数据和运维数据。"),
        ("9.1.2", "样品、任务、机器人、工具、仪器、工位、程序、方法和事件应采用唯一编码。编码规则在同一项目内应保持稳定。"),
        ("9.1.3", "机器人和试验仪器应使用统一时钟源。一般业务数据时间同步误差不应大于1 s；动态协同控制的时间同步误差不应大于一个采样周期，且宜不大于100 ms。"),
        ("9.1.4", "原始动作和仪器数据应只追加、不覆盖。通信中断后补传的数据应保持原时间戳并标记补传状态。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "9.2 接口要求", 2)
    for num, text in [
        ("9.2.1", "系统应提供与试验室信息管理、委托管理、仪器数据采集、样品库、环境监控和项目质量管理系统交换数据的接口。"),
        ("9.2.2", "接口文件应说明字段、类型、单位、精度、时间格式、编码、调用方式、权限、异常码、重试规则和版本。"),
        ("9.2.3", "接口写入失败应记录原因并自动重试。重试不得造成重复任务、重复试验、重复数据或关联覆盖。"),
        ("9.2.4", "涉及机器人运动和危险设备启停的控制命令应经过身份认证、状态核验和权限检查。普通业务接口不得直接绕过安全控制。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "9.3 一样一档", 2)
    for num, text in [
        ("9.3.1", "应建立一样一档，至少包括委托与样品信息、分样制样关系、流转与养护、任务配方、机器人及工具、试验仪器、原始数据、异常与人工接管、结果审核、留样和处置记录。"),
        ("9.3.2", "一样一档应支持由报告结果追溯至样品、任务、仪器原始数据和作业过程，并支持导出为可长期读取的格式。"),
        ("9.3.3", "样品、任务或结果关联发生变更时，应进行权限审核并保留变更前后关系。"),
        ("9.3.4", "关键数据完整率应达到100%，一般数据完整率不应低于98%。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "10 异常处置与人工接管", 1)
    add_heading(doc, "10.1 异常分类", 2)
    add_clause(doc, "10.1.1", "异常宜分为样品异常、作业异常、仪器异常、工具异常、环境异常、数据异常和安全异常。")
    add_clause(doc, "10.1.2", "异常状态应根据影响和可恢复性分为自动恢复、暂停等待、人工接管和终止任务。安全异常、样品身份异常和可能影响试验结果的异常不得静默自动恢复。")
    add_clause(doc, "10.1.3", "异常记录应至少包含样品、任务、工位、时间、设备、当前步骤、状态参数、证据、处置动作和结果。")

    add_heading(doc, "10.2 自动恢复", 2)
    for num, text in [
        ("10.2.1", "仅在异常原因明确、样品状态未改变、安全条件满足且恢复规则已验证时，系统方可自动恢复。"),
        ("10.2.2", "自动恢复次数应受限制。同一任务同一异常连续发生2次或累计发生3次时，应停止自动恢复并转入人工接管。"),
        ("10.2.3", "自动恢复前后应重新核验样品身份、机器人位置、工具、仪器和数据关联。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "10.3 人工接管", 2)
    for num, text in [
        ("10.3.1", "人工接管应由授权人员执行。接管前应使机器人和危险设备处于安全状态，并确认样品和任务当前状态。"),
        ("10.3.2", "人工调整样品、工具、设备或数据后，应记录调整内容、原因、人员和时间。恢复自动运行前应重新执行必要的身份、位置和安全核验。"),
        ("10.3.3", "人工接管后继续完成的任务应标识为人机协同任务，试验结果应由授权人员重点复核。"),
        ("10.3.4", "无法确认试样状态、试验条件或仪器数据有效性时，应终止任务并按相应试验方法重新取样或试验。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "11 安全防护", 1)
    for num, text in [
        ("11.1", "机器人系统的风险评估和风险减小应符合GB/T 15706、GB 11291.1和GB 11291.2的规定。风险评估应覆盖正常运行、上下料、清洁、示教、维护、故障恢复和人工接管。"),
        ("11.2", "非协作机器人作业单元应设置固定或活动防护、联锁门、光幕、激光扫描或其他保护装置。人员进入安全工作空间时，危险运动应停止或进入经验证的安全状态。"),
        ("11.3", "急停装置应设置在操作、维护和可能发生危险的位置，并符合GB/T 16754的规定。急停复位不得直接导致机器人或试验设备重新启动。"),
        ("11.4", "安全控制系统应符合GB/T 16855.1的设计原则。安全功能的性能等级应根据风险评估确定。"),
        ("11.5", "自动换装、夹持和真空吸附应设置失效安全措施。能源中断时，应防止工具或载荷坠落。"),
        ("11.6", "压力、抗折和其他破坏性试验应在防护关闭后启动。防护打开、机器人未退出或人员侵入时，试验设备不得加载。"),
        ("11.7", "高温、旋转、搅拌、切割、筛分和清洁设备应设置相应的防烫、防卷入、防夹、防飞溅和防粉尘措施。"),
        ("11.8", "示教、维修和清洁模式应采用受控权限、限速或保持运行装置，并防止其他人员远程启动设备。"),
        ("11.9", "所有关键安全联锁和急停功能的验收测试通过率应达到100%。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "12 安装调试与作业验证", 1)
    add_heading(doc, "12.1 安装与调试", 2)
    for num, text in [
        ("12.1.1", "安装前应核查机器人基础、工作空间、负载、供电、通信、压缩空气、真空、给排水、防护和环境条件。"),
        ("12.1.2", "安装完成后应进行机械连接、电气连接、急停、安全联锁、工具中心点、坐标系、视觉标定、工位位置和接口检查。"),
        ("12.1.3", "手眼标定、工具中心点和关键工位应使用可溯源基准进行确认。机器人、相机、工具、工位或基础发生移动后应重新确认。"),
        ("12.1.4", "调试应依次进行单设备、单动作、单任务、联动任务和连续运行测试，不得直接以全速自动模式进行首次调试。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "12.2 作业验证", 2)
    for num, text in [
        ("12.2.1", "每类正式自动化作业应验证样品适用范围、抓取与放置、关键动作参数、仪器协同、清洁、异常和人工接管。"),
        ("12.2.2", "位置、质量、时间、温度或其他量化操作应采用可溯源方法进行比对。每类作业比对数据不应少于30组；项目规模不足时不应少于10组并覆盖全部典型样品。"),
        ("12.2.3", "涉及试验结果的自动操作应与标准人工操作进行平行试验。每个代表性试验项目的配对样本不应少于10组，总数不应少于30组。"),
        ("12.2.4", "平行试验结果应满足相应试验方法的重复性、再现性或项目批准的比对要求；自动作业不得产生统计显著的系统偏差。"),
        ("12.2.5", "任务配方、机器人程序、工具、仪器方法、样品类型或环境条件发生影响结果的变化后，应重新验证受影响的作业。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "12.3 试运行", 2)
    for num, text in [
        ("12.3.1", "验收前试运行不应少于30 d，并宜覆盖不少于100个自动任务、全部关键作业和全部异常类型；生产任务不足时，应采用受控模拟补充。"),
        ("12.3.2", "连续运行测试宜覆盖一个完整工作班次或连续100个任务。测试期间不得通过删除失败任务提高成功率。"),
        ("12.3.3", "试运行应统计任务成功率、抓取成功率、设备在线率、数据完整率、人工接管、异常恢复和安全联锁情况。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "13 系统检验与验收", 1)
    add_heading(doc, "13.1 一般规定", 2)
    for num, text in [
        ("13.1.1", "系统应依次进行设备进场检验、安装调试检验、功能检验、作业性能检验、平行比对、安全检验、数据检验、试运行和验收。"),
        ("13.1.2", "验收应以批准的建设方案、风险评估、接口文件、试验任务配方、本文件及相关试验方法为依据。"),
        ("13.1.3", "验收样品应覆盖实际使用的主要规格、质量、形状、表面状态和环境条件。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "13.2 检验内容", 2)
    for num, text in [
        ("13.2.1", "设备进场检验应核查型号、负载、工作范围、技术文件、合格证明、校准记录、工具、附件和安全装置。"),
        ("13.2.2", "安装调试检验应核查基础、工位、防护、供电通信、坐标与标定、工具中心点、接口、急停和联锁。"),
        ("13.2.3", "功能检验应覆盖样品识别、任务下发、程序调用、抓取搬运、称量投料、仪器上下料、筛分称量、清洁换样、异常和人工接管。"),
        ("13.2.4", "作业性能检验应评价定位、对中、质量控制、时间控制、抓取成功率、任务成功率和连续运行。"),
        ("13.2.5", "数据检验应覆盖编码、时间同步、原始数据、关联、审计、断网续传、备份恢复和一样一档。"),
        ("13.2.6", "每项关键安全联锁应至少进行3次受控测试；区域入侵、急停、工具锁紧、防护门和危险设备启动条件的测试通过率应达到100%。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "13.3 验收判定", 2)
    for num, text in [
        ("13.3.1", "系统性能应符合表1和附录B的规定。合同、设计文件或试验方法规定严于本文件时，应执行较严要求。"),
        ("13.3.2", "样品身份关联错误、试验方法或程序调用错误、关键安全联锁失效、关键数据缺失、平行比对不合格或可能影响试验结果的异常无法阻断时，验收不应通过。"),
        ("13.3.3", "一般缺陷整改后可复验。复验应覆盖缺陷项及其受影响的关联功能，不得以书面承诺代替实际测试。"),
        ("13.3.4", "验收通过后应形成验收报告、风险评估、问题整改清单、设备与工具清单、任务配方和程序版本清单、接口文件、备份文件及运维移交资料。")
    ]:
        add_clause(doc, num, text)

    add_para(doc, "表1  系统主要验收指标", first_indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, cn_font="黑体", size=10.5)
    add_table(doc,
              ["序号", "指标", "验收要求", "说明"],
              [
                  ["1", "样品—任务—仪器关联准确率", "100%", "抽查不少于100条或全部记录"],
                  ["2", "试验方法和程序调用正确率", "100%", "覆盖全部正式任务配方"],
                  ["3", "工具识别及锁紧确认率", "100%", "自动换装任务"],
                  ["4", "抓取成功率", "≥99%", "按实际抓取次数统计"],
                  ["5", "自动作业任务成功率", "≥98%", "非计划人工干预计为失败"],
                  ["6", "关键设备在线率", "≥98%", "计划停机单列"],
                  ["7", "一般设备在线率", "≥95%", "计划停机单列"],
                  ["8", "关键数据完整率", "100%", "身份、任务、仪器原始数据、异常和审核"],
                  ["9", "一般数据完整率", "≥98%", "按应产生记录统计"],
                  ["10", "关键安全联锁通过率", "100%", "每项不少于3次测试"],
                  ["11", "异常记录完整率", "100%", "含发生、处置、恢复及人员"],
                  ["12", "断网续传", "关键数据不丢失、不重复", "受控断网测试不少于3次"],
                  ["13", "备份恢复", "可恢复且关联完整", "抽取一样一档验证"],
                  ["14", "自动与人工平行比对", "符合相应试验方法", "每项不少于10组，总数不少于30组"]
              ], widths=[1.2, 5.0, 4.0, 6.0], font_size=8.2)

    add_heading(doc, "14 运行维护与数据安全", 1)
    add_heading(doc, "14.1 运行维护", 2)
    for num, text in [
        ("14.1.1", "应建立机器人系统运行维护制度，明确机器人、工具、试验仪器、安全装置、网络、数据和质量业务责任人。"),
        ("14.1.2", "每日运行前应检查安全防护、急停、工具状态、样品识别、仪器校准状态、环境和数据连接。"),
        ("14.1.3", "应按制造商、计量和使用环境要求维护机器人和工具。工具磨损、真空泄漏、夹持力变化、碰撞或拆装后应重新确认。"),
        ("14.1.4", "摄像机、机器人、工具、工位、货架或仪器位置变化后，应检查坐标、标定和程序适用性。"),
        ("14.1.5", "应按月统计任务成功率、抓取失败、人工接管、重复异常、设备离线和数据缺失。连续性能下降时应暂停受影响任务并开展原因分析。"),
        ("14.1.6", "软件、程序、任务配方和接口升级前应进行测试和备份；升级失败时应能够回退至经验证版本。"),
        ("14.1.7", "系统故障期间应启用批准的人工或替代作业。故障恢复后应核查漏采、未完成任务、样品位置和数据关联。")
    ]:
        add_clause(doc, num, text)

    add_heading(doc, "14.2 数据和网络安全", 2)
    for num, text in [
        ("14.2.1", "系统网络安全应符合GB/T 22239的有关规定，并根据系统定级和项目管理要求采取相应防护措施。"),
        ("14.2.2", "机器人控制网络、试验仪器网络和管理网络宜分区部署。远程运维应经过授权并记录全过程日志。"),
        ("14.2.3", "用户权限应按最小授权原则配置。程序发布、任务配方修改、人工接管、数据修正、结果关联和档案导出应使用实名账号并进行审计。"),
        ("14.2.4", "关键数据传输和存储应采取完整性保护措施。样品、人员和项目敏感信息应按法律法规和项目要求进行访问控制。"),
        ("14.2.5", "应制定数据备份、灾难恢复和网络安全事件应急预案，并定期开展恢复验证。")
    ]:
        add_clause(doc, num, text)


def appendices(doc: Document):
    doc.add_page_break()
    add_heading(doc, "附录A（规范性）  典型自动化作业及最低功能", 1)
    add_para(doc, "表A.1规定了公路工程智慧试验室典型机器人自动化作业及最低功能。项目采用表中未列出的作业时，应根据相应试验方法和本文件第12章完成作业验证。")
    add_para(doc, "表A.1  典型自动化作业及最低功能", first_indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, cn_font="黑体", size=10.5)
    rows = [
        ["1", "样品接收与入库", "识读、称重、分配库位、搬运", "身份唯一；库位正确；异常隔离"],
        ["2", "样品拆分与分装", "容器识别、定量分样、封装、编码", "父子关联；质量可核验；无交叉污染"],
        ["3", "天平上下料", "去皮、放置、稳定判定、读数关联", "无偏载碰撞；称量值稳定；数据绑定"],
        ["4", "粉料和颗粒取料", "定量取料、分次补料、残留检查", "符合目标质量和允许偏差"],
        ["5", "自动投料", "顺序控制、实际质量记录、撒漏识别", "保存分次值；异常阻断"],
        ["6", "搅拌与制样辅助", "容器装卸、程序调用、时间控制", "执行相应试验方法；程序正确"],
        ["7", "装模与成型辅助", "试模识别、装料、刮平、振实", "不改变试样状态；参数可追溯"],
        ["8", "试件脱模", "试模定位、脱模、外观检查、搬运", "不得损伤试件；异常隔离"],
        ["9", "养护入库和出库", "龄期核验、库位管理、搬运", "身份、龄期、环境和库位正确"],
        ["10", "烘箱上下料", "耐热抓取、温度核验、计时、冷却流转", "防烫；时间温度符合方法"],
        ["11", "干燥后称量", "冷却状态核验、称量、重复干燥调度", "满足恒重判定和时间要求"],
        ["12", "集料自动筛分", "筛组核验、投料、振筛、完成判定", "筛序100%正确；参数符合方法"],
        ["13", "筛余称量与级配", "分级收集、称量、总质量校核、计算", "筛余关联正确；质量差符合方法"],
        ["14", "混凝土试件抗压上下料", "试件识别、表面清理、对中、装卸", "龄期正确；对中满足方法；防护联锁"],
        ["15", "抗折或其他试件上下料", "方向核验、支点定位、装卸", "放置方向和位置符合试验方法"],
        ["16", "压力试验废样转运", "卸载确认、残片收集、身份核验", "与未试样品隔离；防飞溅"],
        ["17", "工具自动换装", "工具识别、锁紧、能源连接、状态确认", "识别和锁紧确认率100%"],
        ["18", "容器和工具清洁", "清扫、吹扫、冲洗、干燥、完成确认", "清洁配方受控；无交叉污染"],
        ["19", "废液和固废转运", "分类识别、密闭搬运、容器状态确认", "无泄漏、扬尘和误处置"],
        ["20", "自动留样与处置", "期限核验、库位释放、身份复核", "审核状态和期限满足要求"]
    ]
    add_table(doc, ["序号", "作业类别", "最低自动化功能", "关键控制要求"], rows,
              widths=[1.0, 4.0, 6.0, 6.0], font_size=8.0)

    doc.add_page_break()
    add_heading(doc, "附录B（规范性）  核心数据字段和指标计算方法", 1)
    add_heading(doc, "B.1 核心数据字段", 2)
    add_para(doc, "表B.1规定了机器人自动化作业核心数据字段。具体项目可增加字段，但不得删除与样品身份、原始数据、安全、异常和审核有关的必选字段。")
    add_para(doc, "表B.1  核心数据字段", first_indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, cn_font="黑体", size=10.5)
    fields = [
        ["1", "project_id", "项目唯一编码", "字符串", "是"],
        ["2", "laboratory_id", "试验室唯一编码", "字符串", "是"],
        ["3", "sample_id", "样品或试件唯一编码", "字符串", "是"],
        ["4", "parent_sample_id", "父样品编码", "字符串", "条件必选"],
        ["5", "task_id", "试验任务唯一编码", "字符串", "是"],
        ["6", "test_method", "试验方法及版本", "字符串", "是"],
        ["7", "recipe_version", "任务配方版本", "字符串", "是"],
        ["8", "program_version", "机器人程序版本", "字符串", "是"],
        ["9", "robot_id", "机器人唯一编码", "字符串", "是"],
        ["10", "tool_id", "末端执行器或工具编码", "字符串", "是"],
        ["11", "instrument_id", "试验仪器编码", "字符串", "条件必选"],
        ["12", "station_id", "工位或库位编码", "字符串", "是"],
        ["13", "timestamp", "时间及所属时区", "日期时间", "是"],
        ["14", "operation_step", "当前作业步骤", "字符串", "是"],
        ["15", "operation_parameter", "位置、速度、力、质量、时间等参数", "数值/对象", "条件必选"],
        ["16", "raw_data_uri", "仪器原始数据或动作日志引用", "URI", "条件必选"],
        ["17", "environment_data", "温度、湿度等环境数据", "对象", "条件必选"],
        ["18", "quality_flag", "正常、缺失、补传、异常、修正", "枚举", "是"],
        ["19", "operation_result", "完成、失败、人工接管、终止", "枚举", "是"],
        ["20", "event_id", "异常事件唯一编码", "字符串", "异常时必选"],
        ["21", "evidence_uri", "图片、视频、曲线或日志证据", "URI", "异常时必选"],
        ["22", "takeover_record", "人工接管及恢复记录", "字符串/URI", "接管时必选"],
        ["23", "review_status", "结果审核状态", "枚举", "是"],
        ["24", "audit_log", "修改、授权、版本和导出审计", "字符串/URI", "条件必选"]
    ]
    add_table(doc, ["序号", "字段", "含义", "类型", "必选性"], fields,
              widths=[0.9, 3.6, 6.0, 2.8, 2.8], font_size=7.8)

    add_heading(doc, "B.2 指标计算方法", 2)
    add_clause(doc, "B.2.1", "自动作业任务成功率按式（B.1）计算：R_s=N_s/N_t×100%。式中，R_s为任务成功率；N_s为无需非计划人工干预且满足完成条件的任务数；N_t为统计周期内应执行的任务总数。因系统故障取消、重复执行和人工接管完成的任务不得计入成功任务。")
    add_clause(doc, "B.2.2", "抓取成功率按式（B.2）计算：R_g=N_g/N_a×100%。式中，R_g为抓取成功率；N_g为一次抓取后身份、数量和保持状态均正确的次数；N_a为抓取尝试总次数。")
    add_clause(doc, "B.2.3", "设备在线率按式（B.3）计算：R_o=T_o/T_s×100%。式中，R_o为在线率；T_o为设备应运行期间处于正常在线且数据有效的时间；T_s为设备应处于运行或监测状态的时间。计划停机、无任务时段和无效数据判定规则应在验收方案中明确。")
    add_clause(doc, "B.2.4", "数据完整率按式（B.4）计算：R_c=N_v/N_e×100%。式中，R_c为数据完整率；N_v为符合时间、格式、量程和关联要求的有效记录数；N_e为按任务配方和采集规则应产生的记录数。")
    add_clause(doc, "B.2.5", "关联准确率按样品、任务、试验方法、机器人、工具、仪器和结果全部关联正确的记录数占抽查记录总数的比例计算。任一关键关联错误，该条记录应判定为关联错误。")
    add_clause(doc, "B.2.6", "自动与人工平行比对应报告两种作业的平均值、标准差、差值、最大绝对差和相对偏差，并按相应试验方法规定的重复性、再现性或批准的统计方法评价。")


def main():
    doc = Document()
    configure_document(doc)
    for section in doc.sections:
        add_page_number(section.footer.paragraphs[0])

    add_cover(doc)
    add_front_matter(doc)
    add_toc(doc)
    chapter_1_4(doc)
    chapter_5_8(doc)
    chapter_9_14(doc)
    appendices(doc)

    # Set table row repeat for first rows where possible.
    for table in doc.tables:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)

    doc.core_properties.title = "公路工程智慧试验室机器人自动化作业技术规程"
    doc.core_properties.subject = "团体标准征求意见稿初稿"
    doc.core_properties.author = "OpenAI — 按用户需求编制"
    doc.core_properties.keywords = "公路工程, 智慧试验室, 机器人, 自动化作业, 团体标准"

    doc.save(OUTPUT)
    print(f"Generated: {OUTPUT}")


if __name__ == "__main__":
    main()
